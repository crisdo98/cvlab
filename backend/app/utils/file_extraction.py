"""
File Extraction Utilities

Provides utilities for extracting text content from various file formats
including PDF, DOCX, and handling Base64-encoded files.
"""

import base64
import io
import logging
from typing import Optional

logger = logging.getLogger(__name__)


def extract_text_from_base64(base64_content: str, file_format: str) -> str:
    """
    Extract text from Base64-encoded file content.
    
    Args:
        base64_content: Base64-encoded file content
        file_format: File format (pdf, docx, txt, md, etc.)
        
    Returns:
        Extracted text content
        
    Raises:
        ValueError: If file format is not supported or extraction fails
    """
    try:
        # Decode Base64
        file_bytes = base64.b64decode(base64_content)
        
        # Extract based on format
        if file_format.lower() in ['txt', 'md', 'markdown', 'text']:
            return extract_text_from_text(file_bytes)
        elif file_format.lower() == 'pdf':
            return extract_text_from_pdf(file_bytes)
        elif file_format.lower() in ['docx', 'doc']:
            return extract_text_from_docx(file_bytes)
        else:
            # Try as plain text
            logger.warning(f"Unknown format '{file_format}', attempting as plain text")
            return extract_text_from_text(file_bytes)
            
    except Exception as e:
        logger.error(f"Failed to extract text from {file_format}: {e}")
        raise ValueError(f"Failed to extract text from file: {str(e)}")


def extract_text_from_text(file_bytes: bytes) -> str:
    """Extract text from plain text file bytes."""
    try:
        # Try UTF-8 first
        return file_bytes.decode('utf-8')
    except UnicodeDecodeError:
        # Fallback to latin-1
        try:
            return file_bytes.decode('latin-1')
        except Exception as e:
            raise ValueError(f"Failed to decode text file: {e}")


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF file bytes."""
    try:
        import PyPDF2
        
        pdf_file = io.BytesIO(file_bytes)
        reader = PyPDF2.PdfReader(pdf_file)
        
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        
        extracted_text = '\n\n'.join(text_parts)
        
        if not extracted_text.strip():
            raise ValueError("No text could be extracted from PDF (might be scanned/image-based)")
        
        return extracted_text
        
    except ImportError:
        raise ValueError("PyPDF2 library not installed. Cannot extract PDF text.")
    except Exception as e:
        logger.error(f"PDF extraction failed: {e}")
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX file bytes."""
    try:
        from docx import Document
        
        docx_file = io.BytesIO(file_bytes)
        doc = Document(docx_file)
        
        text_parts = []
        
        # Extract paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        
        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        
        extracted_text = '\n\n'.join(text_parts)
        
        if not extracted_text.strip():
            raise ValueError("No text could be extracted from DOCX file")
        
        return extracted_text
        
    except ImportError:
        raise ValueError("python-docx library not installed. Cannot extract DOCX text.")
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")


def is_base64(content: str) -> bool:
    """
    Check if a string is Base64-encoded.
    
    Args:
        content: String to check
        
    Returns:
        True if content appears to be Base64-encoded
    """
    try:
        # Base64 strings should only contain these characters
        if not all(c in 'ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789+/=' for c in content.strip()):
            return False
        
        # Try to decode
        base64.b64decode(content)
        return True
    except Exception:
        return False
